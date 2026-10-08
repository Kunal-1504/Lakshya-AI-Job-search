"""Independent artifact inspection. Fabrications fail even if an LLM approves them."""

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from lakshya.llm import LLM
from lakshya.models import Job, Profile, TailorPlan
from lakshya.resume_tailor import HEADINGS, contact_line, validate_plan
from lakshya.schemas import CriticReport
from lakshya.services.store import artifact_hash
from lakshya.utils import mentions, normalize


def review_bundle(
    profile: Profile, job: Job, files: dict[str, str], llm: LLM | None = None
) -> CriticReport:
    document = Document(files["docx"])
    allowed = {normalize(f.text) for f in profile.facts}
    allowed.update(normalize(v) for v in HEADINGS.values())
    allowed.update({normalize(profile.contact.name or "Resume"), normalize(contact_line(profile))})
    skills = {normalize(s) for s in profile.skills}
    unsupported: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text or normalize(text) in allowed:
            continue
        if all(normalize(s.strip()) in skills for s in text.split(",")):
            continue
        unsupported.append(text)
    issues = []
    pdf = PdfReader(files["pdf"])
    if len(pdf.pages) != 1:
        issues.append("PDF must have exactly one page")
    if document.tables:
        issues.append("ATS resume cannot contain layout tables")
    pdf_text = " ".join(page.extract_text() or "" for page in pdf.pages)
    docx_text = " ".join(p.text for p in document.paragraphs)

    def compact(text: str) -> str:
        return re.sub(r"\s+", "", text).casefold()

    if compact(pdf_text) != compact(docx_text):
        issues.append("PDF and DOCX content differ")
    plan = TailorPlan.model_validate_json(Path(files["plan"]).read_text())
    try:
        validate_plan(profile, plan)
    except ValueError as exc:
        unsupported.append(str(exc))
    missing = [
        s for s in profile.skills if mentions(job.description, s) and not mentions(docx_text, s)
    ]
    typos = []
    if llm:
        critique = llm.structured(
            "Independently inspect resume claims against verbatim master evidence. Flag unsupported "
            "skills, metrics, experience, degrees and typos. Missing keywords cannot justify fabrication. "
            "artifact_sha256 is supplied; return it unchanged.",
            {
                "master_profile": profile.model_dump(mode="json"),
                "resume": docx_text,
                "job": job.model_dump(mode="json"),
                "artifact_sha256": artifact_hash(files),
            },
            CriticReport,
        )
        unsupported.extend(critique.unsupported_claims)
        issues.extend(critique.formatting_issues)
        typos.extend(critique.typos)
        missing.extend(critique.missing_keywords)
        if not critique.passed and not (unsupported or issues or typos):
            issues.append("LLM critic requested human review")
    passed = not (unsupported or issues or typos)
    return CriticReport(
        passed=passed,
        score=max(0, 100 - 30 * len(unsupported) - 15 * len(issues) - 5 * len(typos)),
        unsupported_claims=list(dict.fromkeys(unsupported)),
        missing_keywords=list(dict.fromkeys(missing)),
        formatting_issues=list(dict.fromkeys(issues)),
        typos=typos,
        artifact_sha256=artifact_hash(files),
    )
