"""Conservative PDF/DOCX/text ingestion with a verbatim evidence catalog."""

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from lakshya.models import Contact, Fact, Profile
from lakshya.utils import extract_skills, write_json

HEADINGS = {
    "summary": {"summary", "profile", "objective", "professional summary", "about me"},
    "skills": {"skills", "technical skills", "core skills", "technologies", "skills and tools"},
    "projects": {"projects", "academic projects", "personal projects", "selected projects"},
    "education": {"education", "academic background", "qualifications"},
    "experience": {"experience", "work experience", "professional experience", "internships"},
    "other": {"certifications", "achievements", "awards", "publications", "activities"},
}


def read_resume(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        text = "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    elif suffix == ".docx":
        doc = Document(str(path))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [
                " | ".join(cell.text for cell in row.cells)
                for table in doc.tables
                for row in table.rows
            ]
        )
    elif suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError("Resume must be .pdf, .docx, .txt or .md")
    if len(text.strip()) < 30:
        raise ValueError("Resume contains too little text. Scanned PDFs need OCR before import.")
    return text


def parse_resume(path: Path, destination: Path | None = None) -> Profile:
    text = read_resume(path)
    lines = [re.sub(r"\s+", " ", line).strip(" \t•") for line in text.splitlines() if line.strip()]
    email = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    phone = re.search(r"(?<!\w)(?:\+\d{1,3}[ -]?)?(?:\d[ ()-]?){10,13}(?!\w)", text)
    links = re.findall(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in/|github\.com/)[\w./-]+", text)
    contact = Contact(
        name=lines[0] if len(lines[0]) <= 80 and not re.search(r"[@|:]", lines[0]) else "",
        email=email.group() if email else "",
        phone=phone.group().strip() if phone else "",
    )
    for link in links:
        url = link if link.startswith("http") else "https://" + link
        if "linkedin.com" in url:
            contact.linkedin = url
        elif "github.com" in url:
            contact.github = url
    section = "summary"
    facts: list[Fact] = []
    for i, line in enumerate(lines):
        heading = line.lower().strip(" :-_")
        found = next((key for key, names in HEADINGS.items() if heading in names), None)
        if found:
            section = found
            continue
        if (
            (i == 0 and contact.name)
            or ("@" in line and email)
            or any(link in line for link in links)
        ):
            continue
        if phone and phone.group().strip() == line:
            continue
        facts.append(Fact(id=f"f{len(facts) + 1:03d}", section=section, text=line))
    if not facts:
        raise ValueError("No resume evidence found after contact extraction")
    skills = extract_skills(text)
    # Preserve additional explicitly listed skills verbatim instead of guessing new ones.
    for fact in facts:
        if fact.section == "skills":
            for term in re.split(r"[,;|]", re.sub(r"^[\w /-]+:\s*", "", fact.text)):
                term = term.strip()
                if 1 < len(term) < 60 and term.casefold() not in {s.casefold() for s in skills}:
                    skills.append(term)
    profile = Profile(
        contact=contact,
        skills=skills,
        projects=[f.text for f in facts if f.section == "projects"],
        education=[f.text for f in facts if f.section == "education"],
        experience=[f.text for f in facts if f.section == "experience"],
        keywords=skills.copy(),
        facts=facts,
        raw_text=text,
        source_file=str(path.resolve()),
        strengths=[f.text for f in facts if f.section in {"summary", "projects"}][:5],
        search_queries=[
            f"{role} {skill}"
            for role in ("AI Engineer", "ML Engineer", "Data Analyst")
            for skill in skills[:2]
        ],
    )
    if destination:
        write_json(destination, profile.model_dump(mode="json"))
    return profile
