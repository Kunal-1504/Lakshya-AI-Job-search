"""Portable Excel tracking and conservative follow-up drafts."""

from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from lakshya.config import Settings
from lakshya.db.base import utc_now
from lakshya.db.models import ApplicationRecord, JobRecord


def draft_followups(session: Session, settings: Settings) -> int:
    count = 0
    for application, job in session.execute(
        select(ApplicationRecord, JobRecord).join(
            JobRecord, ApplicationRecord.job_id == JobRecord.id
        )
    ):
        if application.applied_at and application.status in {"applied", "viewed"}:
            application.follow_up_date = application.applied_at + timedelta(
                days=settings.follow_up_days
            )
            if application.follow_up_date <= utc_now() and not application.follow_up_draft:
                application.follow_up_draft = (
                    f"Hello {job.company} hiring team,\n\n"
                    f"I applied for the {job.title} role on {application.applied_at:%d %B %Y}. "
                    "I would appreciate any update on the application. Thank you for your time.\n"
                )
                count += 1
    return count


def export_excel(session: Session, path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Applications"
    sheet.append(
        [
            "Job Title",
            "Company",
            "Location",
            "Match Score",
            "Resume Path",
            "Apply URL",
            "Applied Date",
            "Status",
            "Follow-up Date",
            "Source",
            "Notes",
        ]
    )
    for application, job in session.execute(
        select(ApplicationRecord, JobRecord)
        .join(JobRecord, ApplicationRecord.job_id == JobRecord.id)
        .order_by(JobRecord.created_at.desc())
    ):
        values = [
            job.title,
            job.company,
            job.location,
            job.match_score,
            application.resume_path,
            job.apply_url,
            application.applied_at.isoformat() if application.applied_at else "",
            application.status,
            application.follow_up_date.isoformat() if application.follow_up_date else "",
            job.source,
            application.notes,
        ]
        # Prevent spreadsheet formula injection from scraped/company/user data.
        sheet.append(
            [
                "'" + value
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                else value
                for value in values
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            60, max(15, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
