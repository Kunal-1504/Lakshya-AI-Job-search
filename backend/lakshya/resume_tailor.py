"""LLM-directed, extractive tailoring: every candidate claim comes from master_profile."""

from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Flowable, Paragraph, Spacer

from lakshya.config import Settings
from lakshya.llm import LLM
from lakshya.models import Job, Profile, TailorPlan
from lakshya.utils import fingerprint, mentions, safe_slug, write_json

HEADINGS = {
    "summary": "Summary",
    "skills": "Technical Skills",
    "projects": "Projects",
    "education": "Education",
    "experience": "Experience",
    "other": "Additional Information",
}


def offline_plan(profile: Profile, job: Job) -> TailorPlan:
    relevant = sorted(
        profile.facts,
        key=lambda f: sum(
            mentions(job.description, s) and mentions(f.text, s) for s in profile.skills
        ),
        reverse=True,
    )
    return TailorPlan(
        fact_ids=[f.id for f in relevant[:24]],
        skill_order=sorted(
            profile.skills, key=lambda s: mentions(job.description, s), reverse=True
        ),
        reasons=["Offline mode: emphasize evidence and skills explicitly mentioned in the JD."],
    )


def validate_plan(profile: Profile, plan: TailorPlan) -> TailorPlan:
    facts = {f.id for f in profile.facts}
    if any(key not in facts for key in plan.fact_ids):
        raise ValueError("Tailoring plan invented an evidence ID; rejected")
    if len(set(plan.fact_ids)) != len(plan.fact_ids):
        raise ValueError("Tailoring plan contains duplicate evidence IDs")
    skills = {s.casefold(): s for s in profile.skills}
    if any(s.casefold() not in skills for s in plan.skill_order):
        raise ValueError("Tailoring plan invented a skill; rejected")
    plan.skill_order = list(dict.fromkeys(skills[s.casefold()] for s in plan.skill_order))
    # Include remaining truthful skills so an LLM omission never destroys the master profile.
    plan.skill_order.extend(s for s in profile.skills if s not in plan.skill_order)
    return plan


def pdf_font() -> str:
    if "LakshyaSans" in pdfmetrics.getRegisteredFontNames():
        return "LakshyaSans"
    for path in (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/Library/Fonts/Arial.ttf"),
    ):
        if path.exists():
            pdfmetrics.registerFont(TTFont("LakshyaSans", str(path)))
            return "LakshyaSans"
    return "Helvetica"


def contact_line(profile: Profile) -> str:
    return " | ".join(
        value for key, value in profile.contact.model_dump().items() if key != "name" and value
    )


def blocks(profile: Profile, plan: TailorPlan) -> OrderedDict[str, list[str]]:
    result: OrderedDict[str, list[str]] = OrderedDict()
    # Reorder sections by relevance, preserving original evidence order WITHIN each section.
    # Prefix closure keeps project/employer/degree context with any selected bullet.
    by_id = {f.id: f for f in profile.facts}
    sections = list(dict.fromkeys(by_id[key].section for key in plan.fact_ids))
    if "skills" not in sections:
        sections.insert(0, "skills")
    if profile.education and "education" not in sections:
        sections.append("education")
    for section in sections:
        if section == "skills":
            if plan.skill_order:
                result[section] = [", ".join(plan.skill_order)]
            continue
        facts = [f for f in profile.facts if f.section == section]
        selected = [i for i, f in enumerate(facts) if f.id in plan.fact_ids]
        last = max(selected) if selected else len(facts) - 1
        result[section] = [f.text for f in facts[: last + 1]]
    return result


def make_story(profile: Profile, sections: dict[str, list[str]], size: float) -> list[Flowable]:
    font = pdf_font()
    body = ParagraphStyle(
        "body",
        fontName=font,
        fontSize=size,
        leading=size * 1.23,
        spaceAfter=3,
        alignment=TA_LEFT,
        textColor=colors.black,
    )
    heading = ParagraphStyle(
        "heading", parent=body, fontSize=size + 1, spaceBefore=7, spaceAfter=3, keepWithNext=True
    )
    name = ParagraphStyle("name", parent=body, fontSize=17, leading=20, spaceAfter=4)
    story: list[Flowable] = [Paragraph(escape(profile.contact.name or "Resume"), name)]
    if contact_line(profile):
        story.append(Paragraph(escape(contact_line(profile)), body))
    story.append(Spacer(1, 5))
    for section, lines in sections.items():
        if lines:
            story.append(Paragraph(HEADINGS[section], heading))
            story.extend(Paragraph(escape(line), body) for line in lines)
    return story


def measured_height(story: list[Flowable], width: float) -> float:
    total = 0.0
    for item in story:
        _, height = item.wrap(width, A4[1])
        total += height + item.getSpaceBefore() + item.getSpaceAfter()
    return total


def fit_one_page(
    profile: Profile, sections: OrderedDict[str, list[str]]
) -> tuple[OrderedDict[str, list[str]], float, list[str]]:
    width, height = A4[0] - 90, A4[1] - 90
    omitted: list[str] = []
    for size in (10.5, 10, 9.5):
        if measured_height(make_story(profile, sections, size), width) <= height - 20:
            return sections, size, omitted
    while measured_height(make_story(profile, sections, 9.5), width) > height - 20:
        removable = next(
            (
                key
                for key in reversed(sections)
                if sections[key] and (key not in {"skills", "education"} or len(sections[key]) > 1)
            ),
            None,
        )
        if removable is None:
            raise ValueError(
                "Contact/skills/education are too long for one page; shorten master_profile"
            )
        omitted.append(sections[removable].pop())
    return sections, 9.5, omitted


def render_pdf(path: Path, profile: Profile, sections: dict[str, list[str]], size: float) -> None:
    from weasyprint import HTML

    content = f"<h1>{escape(profile.contact.name or 'Resume')}</h1>"
    content += f"<p>{escape(contact_line(profile))}</p>"
    for key, lines in sections.items():
        if lines:
            content += f"<h2>{HEADINGS[key]}</h2>" + "".join(
                f"<p>{escape(line)}</p>" for line in lines
            )
    style = f"@page {{ size:A4; margin:45pt; }} body {{font-family:DejaVu Sans,sans-serif; font-size:{size}pt; line-height:1.23;}} p{{margin:0 0 3pt}} h1{{font-size:17pt;margin:0 0 4pt}} h2{{font-size:{size + 1}pt; margin:7pt 0 3pt;break-after:avoid}}"
    HTML(
        string=f"<html><head><style>{style}</style></head><body>{content}</body></html>"
    ).write_pdf(path)
    if len(PdfReader(path).pages) != 1:
        path.unlink(missing_ok=True)
        raise ValueError("PDF overflowed one page; generation rejected")


def render_docx(path: Path, profile: Profile, sections: dict[str, list[str]], size: float) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(A4[0]), Pt(A4[1])
    section.top_margin = section.bottom_margin = Inches(0.625)
    section.left_margin = section.right_margin = Inches(0.625)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(size)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(profile.contact.name or "Resume")
    run.bold, run.font.size = True, Pt(17)
    if contact_line(profile):
        document.add_paragraph(contact_line(profile))
    for key, lines in sections.items():
        if not lines:
            continue
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(7)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(HEADINGS[key])
        run.bold, run.font.size = True, Pt(size + 1)
        for line in lines:
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.widow_control = True
    document.save(str(path))


class ResumeTailor:
    def __init__(self, settings: Settings, llm: LLM | None = None) -> None:
        self.settings, self.llm = settings, llm

    def tailor(
        self, profile: Profile, job: Job, feedback: dict[str, object] | None = None
    ) -> dict[str, str]:
        if self.llm:
            plan = self.llm.structured(
                "Tailor a ONE PAGE ATS resume. HARD RULE: never invent skills, employers, degrees, "
                "metrics or projects. Select ONLY fact IDs from master_profile and reorder ONLY skills "
                "the candidate has. Prioritize sections and evidence for this JD. No candidate prose "
                "may be generated. The renderer copies verified evidence verbatim and keeps original "
                "context order inside each section. Reasons explain selections; they are not resume claims. "
                "Mirror JD keywords only if genuinely present in master_profile.",
                {
                    "master_profile": profile.model_dump(mode="json"),
                    "job": job.model_dump(mode="json"),
                    "critic_feedback": feedback or {},
                },
                TailorPlan,
            )
        else:
            plan = offline_plan(profile, job)
        plan = validate_plan(profile, plan)
        selected, size, omitted = fit_one_page(profile, blocks(profile, plan))
        date = datetime.now(ZoneInfo(self.settings.timezone)).strftime("%Y-%m-%d")
        folder = (
            self.settings.output_dir
            / f"{safe_slug(job.company)}_{safe_slug(job.title)}_{date}_{fingerprint(job.id)[:8]}"
        )
        folder.mkdir(parents=True, exist_ok=True)
        pdf, docx = folder / "resume.pdf", folder / "resume.docx"
        render_pdf(pdf, profile, selected, size)
        render_docx(docx, profile, selected, size)
        relevant_skills = [s for s in plan.skill_order if mentions(job.description, s)][:6]
        note = (
            f"Dear {job.company} hiring team,\n\nI am applying for the {job.title} role. "
            + (
                "My background includes " + ", ".join(relevant_skills) + ". "
                if relevant_skills
                else ""
            )
            + "Please find my resume attached. I would welcome the opportunity to discuss how my "
            "existing experience and projects relate to your team's needs.\n\n"
            + (profile.contact.name or "")
            + "\n"
        )
        cover = folder / "cover_note.txt"
        cover.write_text(note, encoding="utf-8")
        resume_text = folder / "resume.txt"
        text = profile.contact.name + "\n" + contact_line(profile) + "\n\n"
        text += "\n\n".join(
            HEADINGS[k] + "\n" + "\n".join(lines) for k, lines in selected.items() if lines
        )
        resume_text.write_text(text + "\n", encoding="utf-8")
        report = folder / "changes_report.md"
        included = [line for lines in selected.values() for line in lines]
        removed = [
            f.text for f in profile.facts if f.section != "skills" and f.text not in included
        ]
        report.write_text(
            f"# Tailoring changes\n\nCompany: {job.company}\nRole: {job.title}\n"
            f"Source: {job.source} — {job.source_url or job.apply_url}\n"
            f"Master profile SHA256: {fingerprint(profile)}\n\n"
            "Candidate claims are copied verbatim from master_profile; no facts were invented. "
            "Sections are ordered for relevance. Original fact order and preceding context are retained "
            "within each section. Skills are reordered to emphasize genuine JD overlap.\n\n"
            "## Selection reasons\n"
            + "\n".join("- " + reason for reason in plan.reasons)
            + "\n\n## Final section order\n"
            + "\n".join("- " + HEADINGS[k] for k in selected if selected[k])
            + "\n\n## Omitted master evidence\n"
            + ("\n".join("- " + value for value in removed) or "None")
            + "\n\n## Page fitting\n"
            + f"PDF verified as one A4 page; body font {size}pt. "
            + f"{len(omitted)} additional lines omitted for fit. DOCX uses the same content; verify pagination in Word.\n",
            encoding="utf-8",
        )
        plan_path = folder / "tailor_plan.json"
        write_json(plan_path, plan.model_dump(mode="json"))
        return {
            "pdf": str(pdf.resolve()),
            "docx": str(docx.resolve()),
            "cover_note": str(cover.resolve()),
            "report": str(report.resolve()),
            "text": str(resume_text.resolve()),
            "plan": str(plan_path.resolve()),
        }
